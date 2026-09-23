"""Leitura e segmentação da planilha de grade (doc 03 §2).

A regra que mais quebra implementações ingênuas: no .xlsx uma célula mesclada
tem valor apenas na superior-esquerda. Expandir as mesclagens é o primeiro passo
de tudo — sem isso, aulas somem e turmas se multiplicam.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.core.texto import slugificar
from app.services.importador.celula import Aula, interpretar

# Abas → curso + turno. "CICLO" não é curso: expande nos cursos reais (RN-03).
CURSO_DA_ABA: dict[str, tuple[str, str]] = {
    "matutino ciclo basico": ("CICLO", "MATUTINO"),
    "educacao fisica matutino": ("EDUCACAO_FISICA", "MATUTINO"),
    "enfermagem matutino": ("ENFERMAGEM", "MATUTINO"),
    "biomedicina matutino": ("BIOMEDICINA", "MATUTINO"),
    "gti matutino": ("GTI", "MATUTINO"),
    "gti noturno": ("GTI", "NOTURNO"),
    "ciclo basico integrado noturn": ("CICLO", "NOTURNO"),
    "biomedicina noturno": ("BIOMEDICINA", "NOTURNO"),
    "enfermagem noturno": ("ENFERMAGEM", "NOTURNO"),
    "educacao fisica noturno": ("EDUCACAO_FISICA", "NOTURNO"),
}
CURSOS_NO_TEXTO = {
    "biomedicina": "BIOMEDICINA",
    "enfermagem": "ENFERMAGEM",
    "educacao fisica": "EDUCACAO_FISICA",
    "gti": "GTI",
}
_PERIODO = re.compile(r"(\d+)\s*[ºo°]?\s*(?:PER\b|Período)", re.IGNORECASE)
# "* Matricular 3, 4, 5 e 6 períodos" é instrução de matrícula, não o período da
# turma — precisa ser ignorada, senão a turma do 4º PER vira 6º.
_INSTRUCAO_MATRICULA = re.compile(r"matricular", re.IGNORECASE)
_ESTAGIO_COL = re.compile(r"est[áa]gio\s+supervisionado", re.IGNORECASE)


@dataclass
class ColunaTurma:
    """Uma coluna (ou grupo mesclado) do cabeçalho."""

    colunas: list[int]
    cabecalho: list[str]
    cursos: list[str]
    turno: str
    periodo_numero: int | None
    modulo: str | None
    rotulo: str
    aba: str
    e_turma: bool = True
    compartilhada: bool = False
    avisos: list[str] = field(default_factory=list)

    @property
    def referencia(self) -> str:
        letras = ",".join(get_column_letter(c) for c in self.colunas)
        return f"{self.aba}!{letras}"


@dataclass
class LinhaGrade:
    """Uma célula de aula interpretada, com sua origem."""

    aba: str
    celula: str
    coluna_turma: ColunaTurma
    dia_semana: str | None
    horario: str | None
    aula: Aula
    duracao_horarios: int = 1


def ler(caminho: str | Path) -> tuple[list[ColunaTurma], list[LinhaGrade]]:
    wb = load_workbook(caminho, data_only=True)
    colunas: list[ColunaTurma] = []
    linhas: list[LinhaGrade] = []

    for ws in wb.worksheets:
        curso_aba, turno = _curso_da_aba(ws.title)
        grade, mesclas, origem, altura = _expandir_mesclas(ws)
        if not grade:
            continue
        max_r = max(r for r, _ in grade)
        max_c = max(c for _, c in grade)

        cabecalho = _linhas_de_cabecalho(grade, max_c)
        primeira_dado = (max(cabecalho) + 1) if cabecalho else 1
        cols = _montar_colunas(
            ws.title, curso_aba, turno, grade, mesclas, cabecalho, max_c
        )
        colunas.extend(cols)
        por_coluna = {c: ct for ct in cols for c in ct.colunas}

        for r in range(primeira_dado, max_r + 1):
            dia = grade.get((r, 1))
            horario = grade.get((r, 2))
            for c in range(3, max_c + 1):
                ct = por_coluna.get(c)
                if ct is None:
                    continue
                # Mescla vertical = UMA aula de N horários: só a célula de
                # origem vira alocação (doc 03 §2.1).
                if (r, c) in origem and origem[(r, c)] != (r, c):
                    continue
                aula = interpretar(grade.get((r, c)))
                if aula is None:
                    continue
                linhas.append(
                    LinhaGrade(
                        aba=ws.title,
                        celula=f"{get_column_letter(c)}{r}",
                        coluna_turma=ct,
                        dia_semana=dia,
                        horario=horario,
                        aula=aula,
                        duracao_horarios=altura.get((r, c), 1),
                    )
                )
    return colunas, linhas


# ─── internos ────────────────────────────────────────────────────────────────
def _curso_da_aba(titulo: str) -> tuple[str, str]:
    chave = slugificar(titulo)
    if chave in CURSO_DA_ABA:
        return CURSO_DA_ABA[chave]
    for k, v in CURSO_DA_ABA.items():  # nome truncado em 31 chars pelo Excel
        if chave.startswith(k) or k.startswith(chave):
            return v
    turno = "NOTURNO" if "noturno" in chave else "MATUTINO"
    return ("INDEFINIDO", turno)


def _expandir_mesclas(ws):
    """doc 03 §2.1 — expandir ANTES de qualquer outra coisa.

    Devolve (grade, mesclas, origem, altura):
      origem[(r,c)] -> célula superior-esquerda da mescla que cobre (r,c)
      altura[(r,c)] -> nº de linhas da mescla, para `duracao_horarios`
    """
    grade: dict[tuple[int, int], str] = {}
    for linha in ws.iter_rows():
        for cel in linha:
            if cel.value is None:
                continue
            texto = re.sub(r"\s+", " ", str(cel.value)).strip()
            if texto and texto != ".":
                grade[(cel.row, cel.column)] = texto

    mesclas = list(ws.merged_cells.ranges)
    origem: dict[tuple[int, int], tuple[int, int]] = {}
    altura: dict[tuple[int, int], int] = {}
    for faixa in mesclas:
        topo = (faixa.min_row, faixa.min_col)
        altura[topo] = faixa.max_row - faixa.min_row + 1
        valor = grade.get(topo)
        for r in range(faixa.min_row, faixa.max_row + 1):
            for c in range(faixa.min_col, faixa.max_col + 1):
                origem[(r, c)] = topo
                if valor:
                    grade.setdefault((r, c), valor)
    return grade, mesclas, origem, altura


def _linhas_de_cabecalho(grade, max_c: int) -> list[int]:
    """Linhas iniciais com A e B vazias e conteúdo a partir de C (doc 03 §2.2)."""
    return [
        r
        for r in range(1, 5)
        if not grade.get((r, 1))
        and not grade.get((r, 2))
        and any(grade.get((r, c)) for c in range(3, max_c + 1))
    ]


def _extrair_periodo(partes: list[str]) -> int | None:
    """O período vem da linha do módulo ("4º PER - Módulo AB", "2 Período -
    Módulo I"), nunca da instrução de matrícula ("Matricular 3, 4, 5 e 6
    períodos") — que traria o maior número em vez do período da turma.
    """
    candidatas = [p for p in partes if not _INSTRUCAO_MATRICULA.search(p)]
    for parte in candidatas or partes:
        m = _PERIODO.search(parte)
        if m:
            return int(m.group(1))
    return None


def _montar_colunas(
    aba, curso_aba, turno, grade, mesclas, cabecalho, max_c
) -> list[ColunaTurma]:
    ultima = max(cabecalho) if cabecalho else 0
    lider: dict[int, int] = {}
    for faixa in mesclas:  # mescla horizontal no ÚLTIMO cabeçalho = mesma turma
        if faixa.min_row == ultima and faixa.max_col > faixa.min_col:
            for c in range(faixa.min_col, faixa.max_col + 1):
                lider[c] = faixa.min_col

    grupos: dict[int, list[int]] = {}
    for c in range(3, max_c + 1):
        if any(grade.get((h, c)) for h in cabecalho) or grade.get((cabecalho[0] if cabecalho else 1, c)):
            grupos.setdefault(lider.get(c, c), []).append(c)

    saida: list[ColunaTurma] = []
    periodo_anterior: int | None = None
    for cols in grupos.values():
        partes = [grade.get((h, cols[0])) for h in cabecalho]
        partes = [p for p in partes if p]
        if not partes:
            continue
        texto = " ".join(partes)
        avisos: list[str] = []

        periodo = _extrair_periodo(partes)
        if periodo is None and periodo_anterior is not None:
            periodo = periodo_anterior  # sub-turma B sem cabeçalho próprio
            avisos.append("PERIODO_HERDADO")
        if periodo is not None:
            periodo_anterior = periodo

        e_turma = not _ESTAGIO_COL.search(texto)
        if not e_turma:
            avisos.append("COLUNA_NAO_E_TURMA")

        if curso_aba == "CICLO":
            alvo = slugificar(texto)
            cursos = sorted({v for k, v in CURSOS_NO_TEXTO.items() if k in alvo})
            compartilhada = True
            if not cursos:
                cursos = ["INDEFINIDO"]
                avisos.append("CURSO_COMPARTILHADO_INDEFINIDO")
            else:
                avisos.append("CICLO_BASICO_EXPANDIDO")
        else:
            cursos, compartilhada = [curso_aba], False

        modulo = None
        alvo_mod = next((p for p in partes if "Módulo" in p), None)
        if alvo_mod:
            modulo = re.sub(r"^.*?Módulo\s*", "", alvo_mod).strip() or None

        rotulo = next(
            (p for p in partes if re.search(r"turma", p, re.IGNORECASE)),
            "Turma Única",
        ).strip()

        saida.append(
            ColunaTurma(
                colunas=cols,
                cabecalho=partes,
                cursos=cursos,
                turno=turno,
                periodo_numero=periodo,
                modulo=modulo,
                rotulo=rotulo,
                aba=aba,
                e_turma=e_turma,
                compartilhada=compartilhada,
                avisos=avisos,
            )
        )
    return saida
